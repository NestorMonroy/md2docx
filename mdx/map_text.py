#!/usr/bin/env python3
"""
Text element mappers: headings, paragraphs, code blocks, page breaks.

This module handles conversion of textual HTML elements to DOCX with:
- Comprehensive error handling
- Input validation
- Fallback behaviors
- Style application with error recovery
"""

import sys
from typing import Dict, Any, Optional

from bs4 import Tag
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_heading(doc, style_config: Dict[str, Any], element) -> None:
    """
    Add heading to document with proper styling.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
        element: HTML heading element (h1-h6)
    """
    # Validate using hasattr instead of isinstance
    if not hasattr(doc, 'add_paragraph'):
        raise TypeError(f"doc must be a valid Document object")

    if not hasattr(element, 'name') or not hasattr(element, 'get_text'):
        raise TypeError(f"element must be a valid Tag object")

    tag_name = element.name.lower() if element.name else None

    if tag_name not in ("h1", "h2", "h3", "h4", "h5", "h6"):
        raise ValueError(f"Invalid heading tag: {tag_name}")

    # Extract text
    text = element.get_text(strip=True)

    if not text:
        return  # Skip empty headings

    # Get heading level
    level = int(tag_name[1])

    # Get style name from config
    try:
        headings = style_config.get("styles", {}).get("headings", {})
        style_name = headings.get(tag_name, f"Heading {level}")
    except Exception as e:
        print(f"WARNING: Error getting heading style: {e}", file=sys.stderr)
        style_name = f"Heading {level}"

    # Add heading with error handling
    try:
        para = doc.add_paragraph(text, style=style_name)
    except KeyError:
        # Fallback if style not found in template
        print(f"WARNING: Style '{style_name}' not found, using fallback", file=sys.stderr)
        para = doc.add_paragraph(text)

        # Apply basic heading formatting
        if para.runs:
            run = para.runs[0]
            run.font.size = Pt(18 - (level * 2))
            run.font.bold = True


def add_paragraph(doc, style_config: Dict[str, Any], element) -> None:
    """
    Add paragraph to document.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
        element: HTML paragraph element
    """
    # Validate using hasattr
    if not hasattr(doc, 'add_paragraph'):
        raise TypeError(f"doc must be a valid Document object")

    if not hasattr(element, 'get_text'):
        raise TypeError(f"element must be a valid Tag object")

    # Extract text
    text = element.get_text()

    if not text.strip():
        return  # Skip empty paragraphs

    # Get style name
    try:
        style_name = style_config.get("styles", {}).get("paragraph", "Normal")
    except Exception as e:
        print(f"WARNING: Error getting paragraph style: {e}", file=sys.stderr)
        style_name = "Normal"

    # Add paragraph with error handling
    try:
        doc.add_paragraph(text, style=style_name)
    except KeyError:
        # Fallback if style not found
        print(f"WARNING: Style '{style_name}' not found, using default", file=sys.stderr)
        doc.add_paragraph(text)
    except Exception as e:
        print(f"WARNING: Error adding paragraph: {e}", file=sys.stderr)


def add_code_block(doc, style_config: Dict[str, Any], element) -> None:
    """
    Add code block to document.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
        element: HTML pre element
    """
    # Validate using hasattr
    if not hasattr(doc, 'add_paragraph'):
        raise TypeError(f"doc must be a valid Document object")

    if not hasattr(element, 'find') or not hasattr(element, 'get_text'):
        raise TypeError(f"element must be a valid Tag object")

    # Extract code text, preferring <code> child if present
    try:
        code_element = element.find("code")
        text = code_element.get_text() if code_element else element.get_text()
    except Exception as e:
        print(f"WARNING: Error extracting code text: {e}", file=sys.stderr)
        text = element.get_text()

    if not text.strip():
        return  # Skip empty code blocks

    # Get style name
    try:
        style_name = style_config.get("styles", {}).get("code_block", "Code")
    except Exception as e:
        print(f"WARNING: Error getting code style: {e}", file=sys.stderr)
        style_name = "Code"

    # Add code block with error handling
    try:
        para = doc.add_paragraph(text, style=style_name)
    except KeyError:
        # Fallback with monospace font
        print(f"WARNING: Style '{style_name}' not found, using fallback", file=sys.stderr)
        para = doc.add_paragraph(text)

        if para.runs:
            run = para.runs[0]
            try:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            except Exception as e:
                print(f"WARNING: Error applying code formatting: {e}", file=sys.stderr)
    except Exception as e:
        print(f"WARNING: Error adding code block: {e}", file=sys.stderr)


def add_blockquote(doc, style_config: Dict[str, Any], element) -> None:
    """
    Add blockquote to document.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
        element: HTML blockquote element
    """
    # Validate using hasattr
    if not hasattr(doc, 'add_paragraph'):
        raise TypeError(f"doc must be a valid Document object")

    if not hasattr(element, 'get_text'):
        raise TypeError(f"element must be a valid Tag object")

    # Extract text
    text = element.get_text(strip=True)

    if not text:
        return  # Skip empty blockquotes

    # Get style name
    try:
        style_name = style_config.get("styles", {}).get("blockquote", "Quote")
    except Exception as e:
        print(f"WARNING: Error getting blockquote style: {e}", file=sys.stderr)
        style_name = "Quote"

    # Add blockquote with error handling
    try:
        para = doc.add_paragraph(text, style=style_name)
    except KeyError:
        # Fallback with italic formatting and indent
        print(f"WARNING: Style '{style_name}' not found, using fallback", file=sys.stderr)
        para = doc.add_paragraph(text)

        if para.runs:
            try:
                para.runs[0].font.italic = True
            except Exception:
                pass

        try:
            para.paragraph_format.left_indent = Pt(36)
        except Exception as e:
            print(f"WARNING: Error applying blockquote indent: {e}", file=sys.stderr)
    except Exception as e:
        print(f"WARNING: Error adding blockquote: {e}", file=sys.stderr)


def add_page_break(doc, style_config: Dict[str, Any]) -> None:
    """
    Add page break to document.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
    """
    # Validate using hasattr
    if not hasattr(doc, 'add_page_break'):
        raise TypeError(f"doc must be a valid Document object")

    # Check behavior setting
    try:
        behaviors = style_config.get("behaviors", {})
        hr_as_page_break = behaviors.get("hr_as_page_break", True)
    except Exception as e:
        print(f"WARNING: Error getting behavior config: {e}", file=sys.stderr)
        hr_as_page_break = True

    try:
        if hr_as_page_break:
            doc.add_page_break()
        else:
            # Add horizontal line as alternative
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("─" * 50)
            try:
                run.font.color.rgb = RGBColor(192, 192, 192)
            except Exception:
                pass  # Ignore color errors
    except Exception as e:
        print(f"WARNING: Error adding page break: {e}", file=sys.stderr)


def add_text(doc, style_config: Dict[str, Any], text: str) -> None:
    """
    Add plain text to document.

    Args:
        doc: python-docx Document
        style_config: Style configuration dictionary
        text: Text content to add
    """
    # Validate using hasattr
    if not hasattr(doc, 'add_paragraph'):
        raise TypeError(f"doc must be a valid Document object")

    if not isinstance(text, str):
        raise TypeError(f"text must be str, got {type(text)}")

    if not text or not text.strip():
        return  # Skip empty text

    # Get style name
    try:
        style_name = style_config.get("styles", {}).get("paragraph", "Normal")
    except Exception as e:
        print(f"WARNING: Error getting text style: {e}", file=sys.stderr)
        style_name = "Normal"

    # Add text with error handling
    try:
        doc.add_paragraph(text, style=style_name)
    except KeyError:
        print(f"WARNING: Style '{style_name}' not found, using default", file=sys.stderr)
        doc.add_paragraph(text)
    except Exception as e:
        print(f"WARNING: Error adding text: {e}", file=sys.stderr)